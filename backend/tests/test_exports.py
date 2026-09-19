from datetime import UTC, datetime
from io import BytesIO

from docx import Document
from openpyxl import load_workbook

from app.services.exports import create_docx, create_xlsx


def test_exports_contain_result_data() -> None:
    results = [
        {
            "study_group_name": "ИВТ-1",
            "vk_user_id": 123,
            "first_name": "Иван",
            "last_name": "Иванов",
            "full_name": "Иванов Иван Иванович",
            "responded": True,
            "text": "Готово",
            "attachments": [],
            "responded_at": datetime(2026, 7, 14, 12, 0, tzinfo=UTC),
            "is_late": False,
        },
        {
            "study_group_name": "ИВТ-2",
            "vk_user_id": 456,
            "first_name": "Пётр",
            "last_name": "Петров",
            "full_name": None,
            "responded": False,
            "text": None,
            "attachments": [],
            "responded_at": None,
            "is_late": None,
        },
    ]

    workbook = load_workbook(BytesIO(create_xlsx("Опрос", results)))
    document = Document(BytesIO(create_docx(results)))

    assert workbook.active["C3"].value == "Иванов Иван Иванович"
    assert [paragraph.text for paragraph in document.paragraphs] == [
        "ИВТ-1",
        "Иванов Иван Иванович",
        "ИВТ-2",
        "-",
    ]
